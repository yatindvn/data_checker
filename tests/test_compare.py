"""Automated test plan for Excel Compare. See TEST_PLAN.md for the case list.

    python -m pytest -v
"""
from __future__ import annotations

import os
import sys
import time
from unittest import mock

import numpy as np
import openpyxl
import pandas as pd
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests"))

import compare_engine as ce  # noqa: E402
import make_test_data as td  # noqa: E402

RED = "FFFFC7CE"


# ----------------------------------------------------------------- fixtures
@pytest.fixture(scope="module")
def b2b(tmp_path_factory):
    """The main B2B pair: file 1 original, file 2 with the planted changes."""
    d = tmp_path_factory.mktemp("b2b")
    base = td.make_base()
    base.loc[base[td.KEY] == td.CASE_ONLY_CHANGE[0], td.CASE_ONLY_CHANGE[1]] = "Kush Scientific Agencies"
    f2 = td.make_file2(base)
    p1, p2 = str(d / "B2B_file1.xlsx"), str(d / "B2B_file2.xlsx")
    td.write_xlsx(base, p1)
    td.write_xlsx(f2, p2)
    return p1, p2, d


def run(p1, p2, s1="Sheet1", s2="Sheet1", **opts):
    return ce.compare(p1, s1, p2, s2, ce.Options(**opts))


def xlsx(tmp_path, name, df, sheet="Sheet1"):
    p = str(tmp_path / name)
    df.to_excel(p, index=False, sheet_name=sheet)
    return p


def diff_set(res):
    return set(zip(res.cell_diffs["Key"], res.cell_diffs["Column"]))


PLANTED = {(k, c) for k, c, _ in td.CELL_CHANGES}
TOL = (td.TOLERANCE_CHANGE[0], td.TOLERANCE_CHANGE[1])


# ================================================================ POSITIVE
def test_P01_identical_files_match(b2b):
    p1, _, _ = b2b
    res = run(p1, p1, key_columns=[td.KEY])
    assert not res.has_issues
    assert res.summary["Rows matched"] == td.N_ROWS


def test_P02_missing_column_detected(b2b):
    res = run(b2b[0], b2b[1], key_columns=[td.KEY])
    hi = res.header_issues
    assert ((hi.Column == td.REMOVED_COLUMN) & (hi.Issue == "Missing in File 2")).any()


def test_P03_extra_column_detected(b2b):
    hi = run(b2b[0], b2b[1], key_columns=[td.KEY]).header_issues
    assert ((hi.Column == td.ADDED_COLUMN) & (hi.Issue == "Extra in File 2")).any()


def test_P04_case_only_header_rename_flagged_with_hint(b2b):
    hi = run(b2b[0], b2b[1], key_columns=[td.KEY]).header_issues
    row = hi[hi.Column == td.RENAMED_CASE[0]]
    assert len(row) == 1 and "Only case differs" in row.Note.iloc[0]


def test_P05_key_match_on_shuffled_rows_finds_exactly_planted_changes(b2b):
    res = run(b2b[0], b2b[1], key_columns=[td.KEY])
    assert diff_set(res) == PLANTED | {TOL}


def test_P06_deleted_rows_reported_only_in_file1(b2b):
    res = run(b2b[0], b2b[1], key_columns=[td.KEY])
    assert sorted(res.only_in_a[td.KEY]) == sorted(td.DELETED_KEYS)


def test_P07_added_rows_reported_only_in_file2(b2b):
    res = run(b2b[0], b2b[1], key_columns=[td.KEY])
    assert sorted(res.only_in_b[td.KEY]) == sorted([td.ADDED_KEY, td.DUPLICATED_KEY])


def test_P08_duplicate_key_reported(b2b):
    dk = run(b2b[0], b2b[1], key_columns=[td.KEY]).duplicate_keys
    assert list(dk.Key) == [td.DUPLICATED_KEY] and list(dk.File) == ["File 2"]


def test_P09_numeric_tolerance_hides_tiny_difference(b2b):
    res = run(b2b[0], b2b[1], key_columns=[td.KEY], numeric_tolerance=0.01)
    assert TOL not in diff_set(res)
    assert ("LDX1-5", "Invoice Amount") in diff_set(res)    # big change still caught


def test_P10_ignore_case_hides_case_only_difference(b2b):
    assert td.CASE_ONLY_CHANGE in diff_set(run(b2b[0], b2b[1], key_columns=[td.KEY]))
    res = run(b2b[0], b2b[1], key_columns=[td.KEY], ignore_case=True)
    assert td.CASE_ONLY_CHANGE not in diff_set(res)


def test_P11_trim_spaces_option(b2b):
    on = run(b2b[0], b2b[1], key_columns=[td.KEY], trim_spaces=True)
    off = run(b2b[0], b2b[1], key_columns=[td.KEY], trim_spaces=False)
    assert td.SPACE_ONLY_CHANGE not in diff_set(on)
    assert td.SPACE_ONLY_CHANGE in diff_set(off)


def test_P12_ignore_columns_option(b2b):
    res = run(b2b[0], b2b[1], key_columns=[td.KEY], ignore_columns=["Invoice Amount"])
    assert ("LDX1-5", "Invoice Amount") not in diff_set(res)


def test_P13_same_number_different_format_not_flagged(tmp_path):
    a = xlsx(tmp_path, "a.xlsx", pd.DataFrame({"id": ["1", "2"], "amt": [10, 2.5]}))
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame({"id": ["1", "2"], "amt": ["10.00", "2.50"]}))
    assert not run(a, b, key_columns=["id"]).has_issues


def test_P14_row_by_row_mode(tmp_path):
    a = xlsx(tmp_path, "a.xlsx", pd.DataFrame({"x": [1, 2, 3], "y": ["a", "b", "c"]}))
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame({"x": [1, 9, 3, 4], "y": ["a", "b", "c", "d"]}))
    res = run(a, b)
    assert list(res.cell_diffs["Row #"]) == [3] and list(res.cell_diffs.Column) == ["x"]
    assert len(res.only_in_b) == 1 and len(res.only_in_a) == 0


def test_P15_csv_vs_xlsx(tmp_path):
    df = pd.DataFrame({"id": ["A", "B"], "v": [1, 2]})
    a = xlsx(tmp_path, "a.xlsx", df)
    b = str(tmp_path / "b.csv")
    df.assign(v=[1, 5]).to_csv(b, index=False)
    res = ce.compare(a, "Sheet1", b, "(csv)", ce.Options(key_columns=["id"]))
    assert diff_set(res) == {("B", "v")}


def test_P16_report_sheets_and_red_highlighting(b2b):
    p1, p2, d = b2b
    res = run(p1, p2, key_columns=[td.KEY])
    out = str(d / "report.xlsx")
    ce.write_report(res, out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "Header Issues", "Data Differences",
                             "Rows Only In File 1", "Rows Only In File 2",
                             "Duplicate Keys", "File 2 Highlighted"]
    # issue tables are red (conditional format) and summary says ISSUES FOUND
    assert len(wb["Data Differences"].conditional_formatting) > 0
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Result"] == "ISSUES FOUND"
    # highlighted copy: the changed cell is red, an unchanged one is not
    ws = wb["File 2 Highlighted"]
    hdr = [c.value for c in ws[1]]
    kcol, acol = hdr.index(td.KEY) + 1, hdr.index("Invoice Amount") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, kcol).value == "LDX1-5":
            assert ws.cell(r, acol).fill.fgColor.rgb == RED
            assert ws.cell(r, kcol).fill.fgColor.rgb != RED
        if ws.cell(r, kcol).value == td.ADDED_KEY:        # whole extra row red
            assert ws.cell(r, kcol).fill.fgColor.rgb == RED
    # extra column header in dark red
    assert ws.cell(1, hdr.index(td.ADDED_COLUMN) + 1).fill.fgColor.rgb == "FFC00000"


def test_P17_multi_column_key(tmp_path):
    a = xlsx(tmp_path, "a.xlsx", pd.DataFrame({"g": ["06X", "06X", "19Y"], "inv": [1, 2, 1], "amt": [5, 6, 7]}))
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame({"g": ["19Y", "06X", "06X"], "inv": [1, 2, 1], "amt": [7, 60, 5]}))
    res = run(a, b, key_columns=["g", "inv"])
    assert diff_set(res) == {("06X | 2", "amt")}


def test_P18_performance_100k_rows(tmp_path):
    n = 100_000
    rng = np.random.default_rng(1)
    df = pd.DataFrame({"id": [f"INV{i}" for i in range(n)]} |
                      {f"c{j}": rng.integers(0, 1000, n) for j in range(19)})
    a = xlsx(tmp_path, "a.xlsx", df)
    df2 = df.copy()
    df2.loc[500, "c3"] = -1
    b = xlsx(tmp_path, "b.xlsx", df2.sample(frac=1, random_state=3))
    t = time.perf_counter()
    res = run(a, b, key_columns=["id"])
    ce.write_report(res, str(tmp_path / "r.xlsx"))
    elapsed = time.perf_counter() - t
    print(f"\n100k x 20 compare + report: {elapsed:.1f}s")
    assert diff_set(res) == {("INV500", "c3")}
    assert elapsed < 60


def test_P19_ui_loads_common_columns(b2b):
    import app
    ui = app.App()
    try:
        with mock.patch.object(app.messagebox, "showerror") as err:
            ui.paths[0].set(b2b[0]); ui.sheets[0].set("Sheet1")
            ui.paths[1].set(b2b[1]); ui.sheets[1].set("Sheet1")
            ui._load_columns()
            assert not err.called
        cols = ui.key_list.get(0, "end")
        assert td.KEY in cols and td.REMOVED_COLUMN not in cols and td.ADDED_COLUMN not in cols
    finally:
        ui.destroy()


# ================================================================ NEGATIVE
def test_N01_file_not_found(tmp_path):
    with pytest.raises(FileNotFoundError):
        run(str(tmp_path / "nope.xlsx"), str(tmp_path / "nope2.xlsx"))


def test_N02_not_an_excel_file(tmp_path):
    bad = tmp_path / "fake.xlsx"
    bad.write_text("this is not excel")
    good = xlsx(tmp_path, "g.xlsx", pd.DataFrame({"a": [1]}))
    with pytest.raises(Exception):
        run(str(bad), good)


def test_N03_key_column_missing_in_one_file(b2b):
    with pytest.raises(ValueError, match="Key column"):
        run(b2b[0], b2b[1], key_columns=[td.REMOVED_COLUMN])


def test_N04_sheet_does_not_exist(b2b):
    with pytest.raises(Exception):
        run(b2b[0], b2b[1], s1="NoSuchSheet")


def test_N05_headers_only_no_rows(tmp_path):
    a = xlsx(tmp_path, "a.xlsx", pd.DataFrame(columns=["id", "v"]))
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame(columns=["id", "v"]))
    res = run(a, b, key_columns=["id"])
    assert not res.has_issues and res.summary["Rows matched"] == 0
    ce.write_report(res, str(tmp_path / "r.xlsx"))


def test_N06_completely_empty_sheet(tmp_path):
    a = str(tmp_path / "empty.xlsx")
    openpyxl.Workbook().save(a)
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame({"id": [1]}))
    res = run(a, b, s1="Sheet")          # openpyxl's default blank sheet name
    assert ((res.header_issues.Column == "id") & (res.header_issues.Issue == "Extra in File 2")).any()
    ce.write_report(res, str(tmp_path / "r.xlsx"))


def test_N07_output_not_writable(b2b, tmp_path):
    res = run(b2b[0], b2b[1], key_columns=[td.KEY])
    target = tmp_path / "locked.xlsx"
    target.mkdir()                      # a folder with that name -> cannot write
    with pytest.raises((PermissionError, IsADirectoryError, OSError)):
        ce.write_report(res, str(target))


def test_N08_ui_invalid_tolerance_rejected(b2b):
    import app
    ui = app.App()
    try:
        ui.paths[0].set(b2b[0]); ui.paths[1].set(b2b[1])
        ui.out_path.set(str(b2b[2] / "ui.xlsx")); ui.tolerance.set("abc")
        with mock.patch.object(app.messagebox, "showwarning") as warn, \
             mock.patch.object(app.threading, "Thread") as th:
            ui._start()
        assert warn.called and "tolerance" in warn.call_args[0][0].lower()
        assert not th.called
    finally:
        ui.destroy()


def test_N09_ui_compare_without_files():
    import app
    ui = app.App()
    try:
        with mock.patch.object(app.messagebox, "showwarning") as warn, \
             mock.patch.object(app.threading, "Thread") as th:
            ui._start()
        assert warn.called and not th.called
    finally:
        ui.destroy()


def test_N10_blank_key_values(tmp_path):
    a = xlsx(tmp_path, "a.xlsx", pd.DataFrame({"id": ["A", None, None], "v": [1, 2, 3]}))
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame({"id": ["A", None, None], "v": [1, 2, 4]}))
    res = run(a, b, key_columns=["id"])
    assert list(res.duplicate_keys.Key) == [""] * 2
    assert len(res.cell_diffs) == 1


def test_N11_no_common_columns(tmp_path):
    a = xlsx(tmp_path, "a.xlsx", pd.DataFrame({"x": [1]}))
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame({"y": [1]}))
    res = run(a, b)
    assert set(res.header_issues.Issue) == {"Missing in File 2", "Extra in File 2"}
    ce.write_report(res, str(tmp_path / "r.xlsx"))


def test_N12_file2_has_no_data_rows(tmp_path):
    a = xlsx(tmp_path, "a.xlsx", pd.DataFrame({"id": ["A", "B"], "v": [1, 2]}))
    b = xlsx(tmp_path, "b.xlsx", pd.DataFrame(columns=["id", "v"]))
    res = run(a, b, key_columns=["id"])
    assert len(res.only_in_a) == 2 and len(res.only_in_b) == 0
