"""Tests for: compare all sheets, mapping files, formulas, plain-English logic.
See TEST_PLAN.md (cases P20+ and N13+).
"""
from __future__ import annotations

import os
import sys
from unittest import mock

import openpyxl
import pandas as pd
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests"))

import ai_logic  # noqa: E402
import compare_engine as ce  # noqa: E402
import formula  # noqa: E402
import make_test_data as td  # noqa: E402


@pytest.fixture(scope="module")
def data(tmp_path_factory):
    d = tmp_path_factory.mktemp("feat")
    base = td.make_base_fixed()
    p = {k: str(d / f"{k}.xlsx") for k in ("f1", "summary", "multi1", "multi2", "map", "map_en")}
    td.write_xlsx(base, p["f1"])
    td.write_xlsx(td.make_summary(base), p["summary"])
    m1, m2 = td.make_multi(base)
    td.write_sheets(m1, p["multi1"])
    td.write_sheets(m2, p["multi2"])
    td.write_mapping(p["map"])
    td.write_mapping(p["map_en"], td.MAP_ROWS + [
        ("Avg Tax (English)", "add igst tax and cgst tax then divide by 2", "Avg IGST+CGST", "", "")])
    p["dir"] = d
    return p


def mapped(data, mapping_path=None, ai=None, **opts):
    m = ce.load_mapping(mapping_path or data["map"])
    return ce.compare(data["f1"], "Sheet1", data["summary"], "Sheet1",
                      ce.Options(**opts), mapping=m, ai=ai)


def diffs(res):
    return set(zip(res.cell_diffs["Key"], res.cell_diffs["Column"]))


class FakeAI:
    """Stands in for Claude: returns a fixed formula and records the call."""
    available = True

    def __init__(self, answer):
        self.answer, self.calls = answer, []

    def to_formula(self, text, columns):
        self.calls.append(text)
        formula.validate(self.answer, columns)
        return self.answer


# ================================================================ POSITIVE
def test_P20_all_sheets_matched_by_name(data):
    res = ce.compare_all_sheets(data["multi1"], data["multi2"], ce.Options(key_columns=[td.KEY]))
    by = {r.name: r for r in res}
    assert set(by) == {"B2B", "B2C", "Credit Notes", "Extra"}
    assert by["B2B"].result is not None and not by["B2B"].result.has_issues
    # "B2C" vs " b2c" matched ignoring case/spaces; planted change found
    k, c, _ = td.MULTI_B2C_CHANGE
    assert diffs(by["B2C"].result) == {(k, c)}
    assert by["Credit Notes"].error == "Sheet missing in File 2"
    assert by["Extra"].error == "Sheet missing in File 1"


def test_P21_all_sheets_report(data):
    res = ce.compare_all_sheets(data["multi1"], data["multi2"], ce.Options(key_columns=[td.KEY]))
    out = str(data["dir"] / "multi_report.xlsx")
    ce.write_report(res, out)
    wb = openpyxl.load_workbook(out)
    assert {"Summary", "Data Differences", "F2 B2B", "F2 B2C"} <= set(wb.sheetnames)
    rows = list(wb["Summary"].iter_rows(values_only=True))
    status = {r[0]: r[1] for r in rows[3:] if r[0]}
    assert status["B2B"] == "MATCH" and status["B2C"] == "ISSUES FOUND"
    assert status["Credit Notes"] == "NOT COMPARED" and status["Overall"] == "ISSUES FOUND"
    dd = [c.value for c in wb["Data Differences"][1]]
    assert dd[0] == "Sheet" and wb["Data Differences"].cell(2, 1).value == "B2C"


def test_P22_mapping_renamed_columns_and_formulas(data):
    res = mapped(data)
    expected = {(k, c) for k, c, _ in td.MAP_ERRORS}
    assert diffs(res) == expected
    assert res.summary["Rows matched"] == td.N_ROWS and not len(res.only_in_a) and not len(res.only_in_b)


def test_P23_mapping_formula_values_are_correct(data):
    res = mapped(data)
    row = res.cell_diffs[(res.cell_diffs.Key == "LDX1-9") & (res.cell_diffs.Column == "Avg Tax")]
    assert row["File 2 Value"].iloc[0] in ("999", "999.0")
    base = td.make_base_fixed().set_index(td.KEY)
    expect = (base.loc["LDX1-9", "Igst Tax"] + base.loc["LDX1-9", "Cgst Tax"]) / 2
    assert float(row["File 1 Value"].iloc[0]) == pytest.approx(expect)


def test_P24_per_row_tolerance(data):
    k, c, _ = td.MAP_WITHIN_TOL
    assert (k, c) not in diffs(mapped(data))          # 0.03 inside row tolerance 0.05
    rows = [r if r[0] != c else (*r[:4], "") for r in td.MAP_ROWS]
    p = str(data["dir"] / "map_notol.xlsx")
    td.write_mapping(p, rows)
    assert (k, c) in diffs(mapped(data, p))           # caught without it


def test_P25_blank_mapping_cell_means_same_name(data):
    used = mapped(data).mapping_used.set_index("Output Column")
    assert used.loc["Invoice Date", "File 1 used"] == "[Invoice Date]"
    assert used.loc["Invoice Date", "File 2 used"] == "[Invoice Date]"


def test_P26_plain_english_converted_by_ai(data):
    ai = FakeAI("([Igst Tax] + [Cgst Tax]) / 2")
    res = mapped(data, data["map_en"], ai=ai)
    assert ai.calls == ["add igst tax and cgst tax then divide by 2"]
    used = res.mapping_used.set_index("Output Column")
    assert used.loc["Avg Tax (English)", "File 1 used"] == "([Igst Tax] + [Cgst Tax]) / 2"
    assert "AI converted" in used.loc["Avg Tax (English)", "Note"]
    # same planted error as the formula row
    assert ("LDX1-9", "Avg Tax (English)") in diffs(res)


def test_P27_ai_translator_calls_claude_correctly(data):
    """Real Translator with the API client mocked: checks request + validation + cache."""
    t = ai_logic.Translator("sk-test")
    fake = mock.MagicMock()
    fake.beta.messages.create.return_value = mock.MagicMock(
        stop_reason="end_turn", content=[mock.MagicMock(type="text", text="([Igst Tax] + [Cgst Tax]) / 2")])
    t._client = fake
    cols = ["Igst Tax", "Cgst Tax"]
    assert t.to_formula("add igst and cgst then halve", cols) == "([Igst Tax] + [Cgst Tax]) / 2"
    t.to_formula("add igst and cgst then halve", cols)            # cached
    assert fake.beta.messages.create.call_count == 1
    kw = fake.beta.messages.create.call_args.kwargs
    assert kw["model"] == "claude-opus-5-5" and "Igst Tax" in kw["messages"][0]["content"]


def test_P28_mapping_template(data, tmp_path):
    a = str(tmp_path / "a.xlsx")
    b = str(tmp_path / "b.xlsx")
    pd.DataFrame({"Invoice Number": [1], "Buyer Name": ["x"], "Only A": [1]}).to_excel(a, index=False)
    pd.DataFrame({"invoice_number": [1], "BUYER NAME": ["x"]}).to_excel(b, index=False)
    out = str(tmp_path / "map.xlsx")
    ce.write_mapping_template(a, "Sheet1", b, "Sheet1", out, ["Invoice Number"])
    m = pd.read_excel(out, sheet_name="Mapping", dtype=str, keep_default_na=False)
    assert list(m["File 2"]) == ["invoice_number", "BUYER NAME", ""]
    assert list(m["Key"]) == ["Y", "", ""]
    assert "How to use" in openpyxl.load_workbook(out).sheetnames


@pytest.mark.parametrize("expr, expected", [
    ("([a] + [b]) / 2", ["1.5", "5", "#ERROR"]),
    ("ROUND([a] * 1.005, 1)", ["1", "4", "#ERROR"]),
    ("ROUND(2.5, 0) & \"|\" & ROUND(-2.5, 0)", ["3|-3"] * 3),          # Excel rounding
    ("IF([a] = 0, \"zero\", [a] * 10)", ["10", "40", "#ERROR"]),
    ("[t] & \"-\" & [a]", ["Pune-1", "DELHI-4", "-abc"]),
    ("[t] = \"pune\"", ["TRUE", "FALSE", "FALSE"]),                     # text = ignores case
    ("[a] / [z]", ["#ERROR"] * 3),                                       # divide by zero
    ("IFERROR([a] / [z], 0)", ["0"] * 3),
    ("[b] + [blank]", ["2", "6", "0"]),                                  # blank counts as 0
    ("SUM([a], [b], 1) - MAX([a], [b])", ["2", "5", "1"]),
    ("0.1 + 0.2", ["0.3"] * 3),
])
def test_P29_formula_engine(expr, expected):
    df = pd.DataFrame({"a": ["1", "4", "abc"], "b": ["2", "6", ""], "t": ["Pune", "DELHI", ""],
                       "z": ["0", "0", "0"], "blank": ["", "", ""]})
    assert list(formula.evaluate(expr, df)) == expected


def test_P30_mapping_with_all_sheets(data, tmp_path):
    m = ce.load_mapping(data["map"])
    p1, p2 = str(tmp_path / "m1.xlsx"), str(tmp_path / "m2.xlsx")
    base = td.make_base_fixed()
    td.write_sheets({"Jan": base.head(40), "Feb": base.iloc[40:80]}, p1)
    s = td.make_summary(base).set_index("Inv No", drop=False)
    td.write_sheets({"Jan": s.loc[base[td.KEY].head(40)], "Feb": s.loc[base[td.KEY].iloc[40:80]]}, p2)
    res = ce.compare_all_sheets(p1, p2, ce.Options(), mapping=m)
    by = {r.name: r.result for r in res}
    assert diffs(by["Jan"]) == {(k, c) for k, c, _ in td.MAP_ERRORS if int(k.split("-")[1]) <= 40}
    assert not by["Feb"].has_issues


# ================================================================ NEGATIVE
def test_N13_mapping_missing_required_columns(tmp_path):
    p = str(tmp_path / "bad.xlsx")
    pd.DataFrame({"Foo": ["x"], "Bar": ["y"]}).to_excel(p, index=False)
    with pytest.raises(ValueError, match="Output Column"):
        ce.load_mapping(p)


def test_N14_mapping_column_not_found(data, tmp_path):
    p = str(tmp_path / "m.xlsx")
    td.write_mapping(p, [("Invoice No", "Invoice Number", "Inv No", "Y", ""),
                         ("Tax", "[Igst Tax] + [No Such Col]", "Total Tax", "", "")])
    with pytest.raises(ValueError, match=r"File 1, 'Tax': Column \[No Such Col\] not found"):
        mapped(data, p)


def test_N15_malicious_formula_rejected(data, tmp_path):
    p = str(tmp_path / "m.xlsx")
    td.write_mapping(p, [("Invoice No", "Invoice Number", "Inv No", "Y", ""),
                         ("X", '=__import__("os").system("calc")', "Total Tax", "", "")])
    with pytest.raises(ValueError):
        mapped(data, p)


def test_N16_plain_english_without_api_key(data):
    with pytest.raises(ValueError, match="API key"):
        mapped(data, data["map_en"], ai=ai_logic.Translator(api_key=""))


def test_N17_ai_returns_invalid_formula(data):
    t = ai_logic.Translator("sk-test")
    t._ask = lambda text, cols: "[Igst Tax] + [Made Up Column]"
    with pytest.raises(ValueError, match="invalid formula"):
        mapped(data, data["map_en"], ai=t)
    t._ask = lambda text, cols: "CANNOT: no matching column"
    with pytest.raises(ValueError, match="could not convert"):
        mapped(data, data["map_en"], ai=t)


def test_N18_all_sheets_one_sheet_fails_others_continue(data):
    res = ce.compare_all_sheets(data["multi1"], data["multi2"], ce.Options(key_columns=["x"]))
    errs = {r.name: r.error for r in res}
    assert "Key column" in errs["B2B"] and "Key column" in errs["B2C"]
    ce.write_report(res, str(data["dir"] / "failed_sheets.xlsx"))     # still writes a report


def test_N19_mapping_duplicate_output(tmp_path):
    p = str(tmp_path / "m.xlsx")
    td.write_mapping(p, [("A", "x", "y", "", ""), ("A", "x", "z", "", "")])
    with pytest.raises(ValueError, match="more than once"):
        ce.load_mapping(p)


def test_N20_mapping_bad_tolerance(tmp_path):
    p = str(tmp_path / "m.xlsx")
    td.write_mapping(p, [("A", "x", "y", "", "abc")])
    with pytest.raises(ValueError, match="not a number"):
        ce.load_mapping(p)


@pytest.mark.parametrize("bad", ["([a] + 1", "[a] +* 2", "ROUND([a]", "[a] [a]", "FOO([a])", "1 < [a] < 3"])
def test_N21_formula_syntax_errors(bad):
    with pytest.raises(formula.FormulaError):
        formula.evaluate(bad, pd.DataFrame({"a": ["1"]}))


def test_N22_no_common_sheets(tmp_path):
    a, b = str(tmp_path / "a.xlsx"), str(tmp_path / "b.xlsx")
    td.write_sheets({"One": pd.DataFrame({"x": [1]})}, a)
    td.write_sheets({"Two": pd.DataFrame({"x": [1]})}, b)
    res = ce.compare_all_sheets(a, b, ce.Options())
    assert [r.error for r in res] == ["Sheet missing in File 2", "Sheet missing in File 1"]
    ce.write_report(res, str(tmp_path / "r.xlsx"))


def test_P31_ui_new_controls(data):
    import app
    ui = app.App()
    try:
        ui.all_sheets.set(True); ui._toggle_all()
        assert all(str(cb.cget("state")) == "disabled" for cb in ui.sheet_boxes)
        ui._set_mapping(data["map"])
        assert str(ui.key_list.cget("state")) == "disabled"
        ui._set_mapping("")
        assert str(ui.key_list.cget("state")) == "normal"
    finally:
        ui.destroy()


def test_N23_ui_bad_mapping_file_rejected(data, tmp_path):
    import app
    bad = str(tmp_path / "bad.xlsx")
    pd.DataFrame({"Foo": ["x"]}).to_excel(bad, index=False)
    ui = app.App()
    try:
        ui.paths[0].set(data["f1"]); ui.paths[1].set(data["summary"])
        ui.out_path.set(str(tmp_path / "r.xlsx")); ui._set_mapping(bad)
        with mock.patch.object(app.messagebox, "showwarning") as warn, \
             mock.patch.object(app.threading, "Thread") as th:
            ui._start()
        assert warn.called and "Mapping" in warn.call_args[0][0] and not th.called
    finally:
        ui.destroy()
